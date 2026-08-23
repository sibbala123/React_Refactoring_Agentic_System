"""
Validate the ReactRefactor LangGraph pipeline against real developer commits.

For each (dataset row x mapped smell x changed front-end file), rewind the
target repo to the commit *before* the developer's fix, run our pipeline's
run_task() on that file for that smell, and record whether the pipeline
picked a matching tactic and produced a verified fix (tsc + smell-resolved).

See experiments/results/ for outputs: summary.csv, run_log.txt, and one
folder per task containing before/after file content, our diff, the
developer's diff, and the pipeline's full decision trail.
"""

from __future__ import annotations

import argparse
import csv
import difflib
import json
import re
import subprocess
import sys
import traceback
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

# ── Path bootstrap ──────────────────────────────────────────────────────────
EXPERIMENTS_DIR = Path(__file__).resolve().parent
REPO_ROOT = EXPERIMENTS_DIR.parent  # "7000 project"
REPOS_DIR = EXPERIMENTS_DIR / "_repos"
RESULTS_DIR = EXPERIMENTS_DIR / "results"
XLSX_PATH = EXPERIMENTS_DIR / "react_smell_commits_mapped.xlsx"

if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from agentic_refactor_system.langgraph_pipeline.runner import run_task  # noqa: E402
from agentic_refactor_system.langgraph_pipeline.ingestion.gather_context import (  # noqa: E402
    gather_context_for_smell,
)
from agentic_refactor_system.utils.manifest_utils import stable_task_id  # noqa: E402
from agentic_refactor_system.utils.paths import sanitize_name  # noqa: E402
from agentic_refactor_system.utils.openai_key import ensure_openai_api_key  # noqa: E402

# ── Config ───────────────────────────────────────────────────────────────────
FRONTEND_GLOBS = ["*.jsx", "*.tsx", "*.js"]
EXCLUDE_PATH_PATTERNS = [
    re.compile(p, re.IGNORECASE)
    for p in [
        r"\.test\.",
        r"\.spec\.",
        r"\.stories\.",
        r"/__tests__/",
        r"/__mocks__/",
        r"/cypress/",
        r"/e2e/",
        r"(^|/)tests?/",
    ]
]
STOPWORDS = {"to", "a", "the", "with", "instead", "of", "in", "and", "or", "for", "from", "own"}


# ── Small utilities ──────────────────────────────────────────────────────────

def log(run_log_fh, msg: str) -> None:
    ts = datetime.now(timezone.utc).strftime("%H:%M:%S")
    line = f"[{ts}] {msg}"
    print(line, flush=True)
    run_log_fh.write(line + "\n")
    run_log_fh.flush()


def run_git(args: list[str], cwd: Path, check: bool = True) -> subprocess.CompletedProcess:
    res = subprocess.run(
        ["git", *args],
        cwd=str(cwd),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if check and res.returncode != 0:
        raise RuntimeError(f"git {' '.join(args)} failed in {cwd}:\n{res.stdout}\n{res.stderr}")
    return res


def ensure_clone(owner: str, name: str, run_log_fh) -> Path:
    clone_dir = REPOS_DIR / f"{owner}__{name}"
    if (clone_dir / ".git").exists():
        return clone_dir
    log(run_log_fh, f"cloning {owner}/{name} -> {clone_dir}")
    REPOS_DIR.mkdir(parents=True, exist_ok=True)
    run_git(
        ["clone", "--filter=blob:none", f"https://github.com/{owner}/{name}.git", str(clone_dir)],
        cwd=REPOS_DIR,
    )
    return clone_dir


def reset_and_checkout(clone_dir: Path, sha: str) -> None:
    run_git(["reset", "--hard"], cwd=clone_dir)
    run_git(["clean", "-fd"], cwd=clone_dir)
    run_git(["checkout", "-q", sha], cwd=clone_dir)


def resolve_shas(clone_dir: Path, fix_sha_short: str) -> tuple[str, str]:
    fix_full = run_git(["rev-parse", fix_sha_short], cwd=clone_dir).stdout.strip()
    parent_full = run_git(["rev-parse", f"{fix_sha_short}^1"], cwd=clone_dir).stdout.strip()
    return fix_full, parent_full


def changed_frontend_files(clone_dir: Path, parent_sha: str, fix_sha: str) -> list[str]:
    res = run_git(
        ["diff", "--name-only", parent_sha, fix_sha, "--", *FRONTEND_GLOBS],
        cwd=clone_dir,
    )
    files = [f.strip() for f in res.stdout.splitlines() if f.strip()]
    kept = []
    for f in files:
        if any(p.search(f) for p in EXCLUDE_PATH_PATTERNS):
            continue
        kept.append(f)
    return kept


def developer_diff(clone_dir: Path, parent_sha: str, fix_sha: str, file_path: str) -> str:
    res = run_git(["diff", parent_sha, fix_sha, "--", file_path], cwd=clone_dir, check=False)
    return res.stdout


def normalize_tokens(s: str) -> set[str]:
    return {w for w in re.findall(r"[a-z]+", s.lower()) if w not in STOPWORDS and len(w) > 2}


def tactic_matches_developer(tactic_name: str, tactic_display_name: str, operations_str: str) -> bool:
    if not tactic_name or tactic_name == "NO_TACTIC":
        return False
    tactic_tokens = normalize_tokens(tactic_display_name) | normalize_tokens(tactic_name.replace("_", " "))
    for op in operations_str.split("|"):
        op_tokens = normalize_tokens(op)
        if not op_tokens or not tactic_tokens:
            continue
        overlap = tactic_tokens & op_tokens
        if len(overlap) >= 2 or (overlap and (op_tokens <= tactic_tokens or tactic_tokens <= op_tokens)):
            return True
    return False


def unified_diff_text(before: str, after: str, path_label: str) -> str:
    return "".join(
        difflib.unified_diff(
            before.splitlines(keepends=True),
            after.splitlines(keepends=True),
            fromfile=f"a/{path_label}",
            tofile=f"b/{path_label}",
        )
    )


# ── Task execution ───────────────────────────────────────────────────────────

def build_task_inputs(
    smell_type: str,
    file_path: str,
    clone_dir: Path,
    repo_slug: str,
    project: str,
    commit_url: str,
    developer_operations: str,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]] | None:
    """Return (manifest_task, smell, context) or None if the file doesn't exist at this checkout."""
    abs_path = clone_dir / file_path
    if not abs_path.exists():
        return None

    text = abs_path.read_text(encoding="utf-8", errors="replace")
    num_lines = max(1, len(text.splitlines()))

    smell = {
        "smell_id": f"commit_eval_{stable_task_id(repo_slug, {'smell_type': smell_type, 'file_path': file_path})}",
        "smell_type": smell_type,
        "file_path": file_path,
        "component_name": None,
        "line_start": 1,
        "line_end": num_lines,
        "severity": "medium",
        "confidence": 0.8,
        "detector_metadata": {
            "note": "Smell asserted from developer-commit dataset, not detected by our scanner.",
            "commit_url": commit_url,
            "developer_operations": developer_operations,
        },
    }

    context = gather_context_for_smell(
        smell=smell,
        target_root=clone_dir,
        repo_name=repo_slug,
    )

    task_id = stable_task_id(repo_slug, smell)
    allowed_files = sorted(dict.fromkeys([file_path, *context.get("local_imports", [])]))

    manifest_task = {
        "id": task_id,
        "repo_name": repo_slug,
        "target_root": str(clone_dir),
        "smell_id": smell["smell_id"],
        "smell_type": smell_type,
        "target_file": file_path,
        "symbol_name": None,
        "line_start": 1,
        "line_end": num_lines,
        "allowed_edit_scope": {
            "mode": "bounded_file_and_local_imports",
            "allowed_files": allowed_files,
        },
        "relevant_context_files": context.get("relevant_context_files", [file_path]),
        "build_command": "",
        "validation_commands": [],
        "metadata": {
            "confidence": 0.8,
            "severity": "medium",
            "detector_metadata": smell["detector_metadata"],
        },
    }

    return manifest_task, smell, context


def run_single_validation_task(
    *,
    row_index: int,
    project: str,
    fix_sha: str,
    parent_sha: str,
    commit_url: str,
    smell_type: str,
    developer_operations: str,
    file_path: str,
    clone_dir: Path,
    repo_slug: str,
    run_log_fh,
) -> dict[str, Any]:
    task_folder_name = f"{sanitize_name(file_path)}__{sanitize_name(smell_type)}"
    out_dir = RESULTS_DIR / f"{row_index}_{sanitize_name(project)}_{fix_sha[:8]}" / task_folder_name
    out_dir.mkdir(parents=True, exist_ok=True)

    row_result: dict[str, Any] = {
        "row_index": row_index,
        "project": project,
        "fix_sha": fix_sha,
        "parent_sha": parent_sha,
        "target_file": file_path,
        "mapped_smell": smell_type,
        "developer_operations": developer_operations,
        "pipeline_tactic": "",
        "tactic_matches_developer": False,
        "final_status": "",
        "tsc_pass": "",
        "sniffer_clean": "",
        "retry_count": "",
        "critique_score": "",
        "error": "",
    }

    try:
        reset_and_checkout(clone_dir, parent_sha)

        built = build_task_inputs(
            smell_type, file_path, clone_dir, repo_slug, project, commit_url, developer_operations
        )
        if built is None:
            row_result["final_status"] = "SKIPPED"
            row_result["error"] = "target file does not exist at parent commit"
            log(run_log_fh, f"row {row_index} | {project} | {file_path} | {smell_type} -> SKIPPED (file absent at parent)")
            return row_result

        manifest_task, smell, context = built
        abs_path = clone_dir / file_path
        before_text = abs_path.read_text(encoding="utf-8", errors="replace")

        final_state = run_task(manifest_task, smell, context, show_progress=False)

        after_text = abs_path.read_text(encoding="utf-8", errors="replace") if abs_path.exists() else ""

        status = final_state.get("status", "unknown")
        plan = final_state.get("plan") or {}
        tactic_name = plan.get("tactic_name", "")
        vr = final_state.get("verification_result") or {}
        checks = vr.get("checks", {})
        cr = final_state.get("critique_result") or {}
        error = final_state.get("error")

        row_result.update(
            {
                "pipeline_tactic": tactic_name,
                "tactic_matches_developer": tactic_matches_developer(
                    tactic_name, plan.get("display_name", tactic_name), developer_operations
                ),
                "final_status": status.upper() if isinstance(status, str) else str(status),
                "tsc_pass": checks.get("typecheck", ""),
                "sniffer_clean": checks.get("smell_resolved", ""),
                "retry_count": final_state.get("retry_count", 0),
                "critique_score": cr.get("score", ""),
                "error": (error.splitlines()[0][:200] if error else ""),
            }
        )

        # ── Persist artifacts ────────────────────────────────────────────────
        ext = Path(file_path).suffix or ".txt"
        (out_dir / f"before{ext}").write_text(before_text, encoding="utf-8")
        (out_dir / f"after{ext}").write_text(after_text, encoding="utf-8")
        (out_dir / "pipeline.diff").write_text(
            unified_diff_text(before_text, after_text, file_path), encoding="utf-8"
        )
        (out_dir / "task_summary.json").write_text(
            json.dumps(
                {
                    "task_id": manifest_task["id"],
                    "repo_name": repo_slug,
                    "project": project,
                    "fix_sha": fix_sha,
                    "parent_sha": parent_sha,
                    "target_file": file_path,
                    "smell_type": smell_type,
                    "status": status,
                    "skip_reason": final_state.get("skip_reason"),
                    "error": error,
                    "actionability": final_state.get("actionability"),
                    "plan": plan,
                    "changed_files": final_state.get("changed_files"),
                    "critique_result": cr,
                    "verification_result": vr,
                    "retry_count": final_state.get("retry_count", 0),
                },
                indent=2,
                default=str,
            ),
            encoding="utf-8",
        )
        dev_diff = developer_diff(clone_dir, parent_sha, fix_sha, file_path)
        (out_dir / "developer.diff").write_text(dev_diff, encoding="utf-8")

        log(
            run_log_fh,
            f"row {row_index} | {project} | {file_path} | smell={smell_type} -> "
            f"status={row_result['final_status']} tactic={tactic_name or '-'} "
            f"tsc={row_result['tsc_pass']} sniffer={row_result['sniffer_clean']} "
            f"retries={row_result['retry_count']} critique={row_result['critique_score']}"
            + (f" | ERROR: {row_result['error']}" if row_result["error"] else ""),
        )

    except Exception as exc:  # noqa: BLE001 - never let one bad task stop the run
        row_result["final_status"] = "FAILED"
        row_result["error"] = f"{type(exc).__name__}: {exc}"
        log(run_log_fh, f"row {row_index} | {project} | {file_path} | smell={smell_type} -> CRASHED: {exc}")
        (out_dir / "crash_traceback.txt").write_text(traceback.format_exc(), encoding="utf-8")
    finally:
        try:
            reset_and_checkout(clone_dir, parent_sha)
        except Exception:
            pass

    return row_result


# ── Dataset loading ──────────────────────────────────────────────────────────

def load_rows(limit_rows: int) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["Commits (mapped)"]
    headers = [c.value for c in ws[1]]
    rows = []
    for i, raw in enumerate(ws.iter_rows(min_row=2, max_row=1 + limit_rows, values_only=True), start=1):
        record = dict(zip(headers, raw))
        rows.append(record)
    return rows


def load_done_keys(summary_path: Path) -> set[tuple[int, str, str]]:
    """(row_index, target_file, mapped_smell) triples already present in summary.csv - for resume."""
    if not summary_path.exists():
        return set()
    done = set()
    with open(summary_path, "r", newline="", encoding="utf-8") as fh:
        for rec in csv.DictReader(fh):
            try:
                done.add((int(rec["row_index"]), rec["target_file"], rec["mapped_smell"]))
            except (KeyError, ValueError):
                continue
    return done


# ── Main ─────────────────────────────────────────────────────────────────────

def main() -> None:
    ensure_openai_api_key()

    parser = argparse.ArgumentParser()
    parser.add_argument("--rows", type=int, default=10)
    parser.add_argument("--dry-run", action="store_true", help="Only clone/checkout/list files, no pipeline calls")
    parser.add_argument("--only-row", type=int, default=0, help="If set, only process this 1-indexed row")
    parser.add_argument("--max-tasks", type=int, default=0, help="Stop after this many tasks (0 = no limit)")
    args = parser.parse_args()

    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    REPOS_DIR.mkdir(parents=True, exist_ok=True)

    run_log_path = RESULTS_DIR / "run_log.txt"
    summary_path = RESULTS_DIR / "summary.csv"

    fieldnames = [
        "row_index", "project", "fix_sha", "parent_sha", "target_file", "mapped_smell",
        "developer_operations", "pipeline_tactic", "tactic_matches_developer", "final_status",
        "tsc_pass", "sniffer_clean", "retry_count", "critique_score", "error",
    ]
    write_header = not summary_path.exists()
    summary_fh = open(summary_path, "a", newline="", encoding="utf-8")
    csv_writer = csv.DictWriter(summary_fh, fieldnames=fieldnames)
    if write_header:
        csv_writer.writeheader()
        summary_fh.flush()

    rows = load_rows(args.rows)
    done_keys = load_done_keys(summary_path)

    task_count = 0
    with open(run_log_path, "a", encoding="utf-8") as run_log_fh:
        log(run_log_fh, f"=== run started: rows=1..{len(rows)} dry_run={args.dry_run} only_row={args.only_row} resume_skip={len(done_keys)} ===")

        for row in rows:
            row_index = row["#"]
            if args.only_row and row_index != args.only_row:
                continue

            project = row["Project"]
            fix_sha_short = row["Commit SHA"]
            commit_url = row["Commit URL"]
            developer_operations = row["Refactoring Operation(s)"] or ""
            mapped_smells = [s.strip() for s in (row["Mapped Code Smell(s)"] or "").split(",") if s.strip()]

            owner, name = project.split("/", 1)
            repo_slug = f"{owner}__{name}"

            try:
                clone_dir = ensure_clone(owner, name, run_log_fh)
                fix_full, parent_full = resolve_shas(clone_dir, fix_sha_short)
                reset_and_checkout(clone_dir, parent_full)
                files = changed_frontend_files(clone_dir, parent_full, fix_full)
            except Exception as exc:  # noqa: BLE001
                log(run_log_fh, f"row {row_index} | {project} | FAILED to prep repo: {exc}")
                continue

            log(
                run_log_fh,
                f"row {row_index} | {project} | fix={fix_full[:10]} parent={parent_full[:10]} "
                f"smells={mapped_smells} files={files}",
            )

            if args.dry_run:
                continue

            for smell_type in mapped_smells:
                for file_path in files:
                    if args.max_tasks and task_count >= args.max_tasks:
                        log(run_log_fh, f"=== reached --max-tasks={args.max_tasks}, stopping ===")
                        summary_fh.close()
                        return

                    key = (row_index, file_path, smell_type)
                    if key in done_keys:
                        log(run_log_fh, f"row {row_index} | {project} | {file_path} | {smell_type} -> already done, skipping (resume)")
                        continue

                    result = run_single_validation_task(
                        row_index=row_index,
                        project=project,
                        fix_sha=fix_full,
                        parent_sha=parent_full,
                        commit_url=commit_url,
                        smell_type=smell_type,
                        developer_operations=developer_operations,
                        file_path=file_path,
                        clone_dir=clone_dir,
                        repo_slug=repo_slug,
                        run_log_fh=run_log_fh,
                    )
                    csv_writer.writerow(result)
                    summary_fh.flush()
                    task_count += 1

        log(run_log_fh, f"=== run finished: {task_count} tasks executed ===")

    summary_fh.close()


if __name__ == "__main__":
    main()
